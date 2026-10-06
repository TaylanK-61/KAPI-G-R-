import os
import datetime
import urllib.request
import xml.etree.ElementTree as ET
import flet as ft
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_PATH = "/sdcard/Download/Kapi_Giris_Ve_Fatura_Takip.xlsx"

def get_tcmb_usd_rate(date_str):
    try:
        dt = datetime.datetime.strptime(date_str, "%d.%m.%Y")
    except ValueError:
        return None, "Tarih formatı hatalı (GG.AA.YYYY)"

    for _ in range(7):
        year_str = dt.strftime("%Y%m")
        day_str = dt.strftime("%d%m%Y")
        url = f"https://www.tcmb.gov.tr/kurlar/{year_str}/{day_str}.xml"
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=4) as resp:
                root = ET.fromstring(resp.read())
                for currency in root.findall('Currency'):
                    if currency.get('CurrencyCode') == 'USD':
                        val = currency.find('ForexSelling').text
                        if val:
                            return float(val), dt.strftime("%d.%m.%Y")
        except Exception:
            dt -= datetime.timedelta(days=1)

    try:
        url_today = "https://www.tcmb.gov.tr/kurlar/today.xml"
        req = urllib.request.Request(url_today, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=4) as resp:
            root = ET.fromstring(resp.read())
            for currency in root.findall('Currency'):
                if currency.get('CurrencyCode') == 'USD':
                    return float(currency.find('ForexSelling').text), "Güncel Kuru"
    except Exception:
        pass

    return None, "Kur çekilemedi"

def setup_excel():
    if os.path.exists(EXCEL_PATH):
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kapı Giriş Tablosu"
    ws.views.sheetView[0].showGridLines = True

    ws.merge_cells("A1:K1")
    ws["A1"] = "KAPI GİRİŞ VE FATURA TAKİP TABLOSU"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        "Sıra No", "Tarih", "Fatura Edilecek Firma", "Firma Kodu",
        "Açıklama 1", "Açıklama 2", "Tutar (USD)", "TCMB Satış Kuru",
        "Tutar (TL)", "Fatura Durumu", "Fatura Numarası"
    ]

    ws.row_dimensions[3].height = 25
    for col_i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col_i, value=h)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    widths = [10, 14, 28, 15, 22, 22, 16, 18, 18, 18, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(EXCEL_PATH)

def main(page: ft.Page):
    setup_excel()
    page.title = "Kapı Giriş & TCMB Otomasyonu"
    page.padding = 20
    page.scroll = ft.ScrollMode.AUTO

    bugun = datetime.datetime.now().strftime("%d.%m.%Y")

    txt_tarih = ft.TextField(label="Tarih (GG.AA.YYYY)", value=bugun)
    txt_firma = ft.TextField(label="Fatura Edilecek Firma")
    txt_firma_kodu = ft.TextField(label="Firma Kodu")
    txt_aciklama1 = ft.TextField(label="Açıklama 1")
    txt_aciklama2 = ft.TextField(label="Açıklama 2")
    txt_usd = ft.TextField(label="Tutar (USD)", keyboard_type=ft.KeyboardType.NUMBER)
    dd_durum = ft.Dropdown(
        label="Fatura Durumu",
        value="Bekliyor",
        options=[ft.dropdown.Option("Bekliyor"), ft.dropdown.Option("Faturalandı")]
    )
    txt_fatura_no = ft.TextField(label="Fatura Numarası")

    lbl_result = ft.Text(value="", color=ft.colors.GREEN_700, weight=ft.FontWeight.BOLD)

    def kaydet_click(e):
        tarih = txt_tarih.value.strip()
        firma = txt_firma.value.strip()
        firma_kodu = txt_firma_kodu.value.strip()
        aciklama1 = txt_aciklama1.value.strip()
        aciklama2 = txt_aciklama2.value.strip()
        usd_str = txt_usd.value.strip()
        durum = dd_durum.value
        fatura_no = txt_fatura_no.value.strip()

        if not firma or not usd_str:
            lbl_result.value = "HATA: Firma ve Tutar alanları zorunludur!"
            lbl_result.color = ft.colors.RED_700
            page.update()
            return

        try:
            usd_val = float(usd_str.replace(",", "."))
        except ValueError:
            lbl_result.value = "HATA: Tutar sayısal olmalıdır!"
            lbl_result.color = ft.colors.RED_700
            page.update()
            return

        rate, info = get_tcmb_usd_rate(tarih)
        if rate is None:
            lbl_result.value = f"HATA: Kur çekilemedi ({info})"
            lbl_result.color = ft.colors.RED_700
            page.update()
            return

        tl_val = usd_val * rate

        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb["Kapı Giriş Tablosu"]
        next_row = ws.max_row + 1

        row_data = [next_row - 3, tarih, firma, firma_kodu, aciklama1, aciklama2, usd_val, rate, tl_val, durum, fatura_no]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            if col_idx == 7:
                cell.number_format = '#,##0.00 "$"'
            elif col_idx in (8, 9):
                cell.number_format = '#,##0.00 "TL"'

        wb.save(EXCEL_PATH)

        lbl_result.value = f"KAYIT BAŞARILI!\n1 USD = {rate:.4f} TL\nToplam: {tl_val:,.2f} TL"
        lbl_result.color = ft.colors.GREEN_700

        txt_firma.value = ""
        txt_firma_kodu.value = ""
        txt_aciklama1.value = ""
        txt_aciklama2.value = ""
        txt_usd.value = ""
        txt_fatura_no.value = ""
        page.update()

    btn_save = ft.ElevatedButton(
        text="EXCEL'E KAYDET",
        on_click=kaydet_click,
        style=ft.ButtonStyle(color=ft.colors.WHITE, bg=ft.colors.BLUE_700)
    )

    page.add(
        ft.Text("KAPI GİRİŞ FORMU", size=22, weight=ft.FontWeight.BOLD),
        txt_tarih, txt_firma, txt_firma_kodu,
        txt_aciklama1, txt_aciklama2, txt_usd,
        dd_durum, txt_fatura_no,
        btn_save, lbl_result
    )

ft.app(target=main)
